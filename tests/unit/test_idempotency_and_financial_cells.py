from openpyxl import Workbook

from pharma_data.cleaning import DataCleanAgent
from pharma_data.contracts import DocumentType, ElementType, ParsedDocument, TableCell
from pharma_data.entity_extraction import EntityExtractAgent, PatternExtractor
from pharma_data.parsers.common import make_element
from pharma_data.parsers.structured import SpreadsheetParser
from pharma_data.relation_extraction import RelationExtractAgent


def test_pipeline_artifact_ids_are_stable_for_same_version_and_content() -> None:
    first = make_element(
        document_version_id="version",
        element_type=ElementType.PARAGRAPH,
        reading_order=0,
        text="Hengrui Pharma EPS 1.20",
        parser_name="test",
        parser_version="1",
    )
    second = make_element(
        document_version_id="version",
        element_type=ElementType.PARAGRAPH,
        reading_order=0,
        text="Hengrui Pharma EPS 1.20",
        parser_name="test",
        parser_version="1",
    )
    assert first.element_id == second.element_id

    document = ParsedDocument(
        document_id="document",
        document_version_id="version",
        document_type=DocumentType.FINANCIAL_REPORT,
        elements=[first],
    )
    agent = EntityExtractAgent([PatternExtractor()])
    first_mentions = agent.extract(document)
    second_mentions = agent.extract(document)
    assert [item.mention_id for item in first_mentions] == [
        item.mention_id for item in second_mentions
    ]

    first_assertions = DataCleanAgent().clean(
        document_version_id="version",
        mentions=first_mentions,
        assertions=RelationExtractAgent().extract(document, first_mentions),
    )
    second_assertions = DataCleanAgent().clean(
        document_version_id="version",
        mentions=second_mentions,
        assertions=RelationExtractAgent().extract(document, second_mentions),
    )
    assert [item.assertion_id for item in first_assertions.assertions] == [
        item.assertion_id for item in second_assertions.assertions
    ]


def test_financial_spreadsheet_emits_cell_level_evidence(tmp_path) -> None:
    path = tmp_path / "finance.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Income Statement"
    sheet.append(["Metric", "2025"])
    sheet.append(["Revenue", 100])
    workbook.save(path)

    parsed = SpreadsheetParser().parse(
        path,
        document_id="document",
        document_version_id="version",
        document_type=DocumentType.FINANCIAL_REPORT,
        artifact_id="artifact",
    )
    cell = parsed.elements[0].structured_payload["financial_cells"][0]

    assert cell["statement_type"] == "income_statement"
    assert cell["row_label"] == "Revenue"
    assert cell["column_label"] == "2025"
    assert cell["numeric_value"] == "100"
    assert cell["period_end"] == "2025-12-31"
    assert cell["cell_reference"] == "B2"
    assert cell["evidence_id"]


def test_pdf_financial_table_emits_period_aware_cell_assertions() -> None:
    element = make_element(
        document_version_id="version",
        element_type=ElementType.TABLE,
        reading_order=0,
        text="项目\t本报告期\t上年同期\n营业收入\t1,234.50\t1,000.00",
        parser_name="mineru",
        parser_version="3.4.0",
    ).model_copy(
        update={
            "table_cells": [
                TableCell(row_index=0, column_index=0, text="项目"),
                TableCell(row_index=0, column_index=1, text="本报告期"),
                TableCell(row_index=0, column_index=2, text="上年同期"),
                TableCell(row_index=1, column_index=0, text="营业收入"),
                TableCell(
                    row_index=1,
                    column_index=1,
                    text="1,234.50",
                    numeric_value="1234.50",
                ),
                TableCell(
                    row_index=1,
                    column_index=2,
                    text="1,000.00",
                    numeric_value="1000.00",
                ),
            ]
        }
    )
    document = ParsedDocument(
        document_id="document",
        document_version_id="version",
        document_type=DocumentType.FINANCIAL_REPORT,
        metadata={"issuer": "测试医药股份有限公司"},
        elements=[element],
    )
    mentions = EntityExtractAgent([PatternExtractor()]).extract(document)
    assertions = RelationExtractAgent().extract(document, mentions)

    assert len(assertions) == 2
    assert {item.evidence_table_cell for item in assertions} == {(1, 1), (1, 2)}
    assert {item.qualifiers["period_semantics"] for item in assertions} == {
        "current_period",
        "prior_year_same_period",
    }
    assert all(item.extraction_method == "schema_rule:FINANCIAL_TABLE_CELL" for item in assertions)


def test_numeric_cells_above_value_are_never_treated_as_period_headers() -> None:
    cells = [
        TableCell(row_index=0, column_index=1, text="-80,543,216.59"),
        TableCell(
            row_index=1,
            column_index=1,
            text="1,876,441,961.32",
            numeric_value="1876441961.32",
            header_path=["-80,543,216.59"],
        ),
    ]

    assert RelationExtractAgent._column_header(cells, cells[1]) == ""
    assert RelationExtractAgent._period_details("2026年1-3月")["period_semantics"] == (
        "explicit_period"
    )


def test_research_investment_scopes_are_distinct_metrics() -> None:
    assert RelationExtractAgent._canonical_metric("累计研发投入") == "累计研发投入"
    assert RelationExtractAgent._canonical_metric("费用化研发投入") == "费用化研发投入"
    assert RelationExtractAgent._canonical_metric("被合并方实现的净利润") == (
        "被合并方净利润"
    )


def test_table_values_with_unmodelled_same_period_dimension_are_withheld() -> None:
    element = make_element(
        document_version_id="version",
        element_type=ElementType.TABLE,
        reading_order=0,
        text="项目\t期末余额\t期末余额\n应收账款\t100\t200",
        parser_name="mineru",
        parser_version="3.4.0",
    ).model_copy(
        update={
            "table_cells": [
                TableCell(row_index=0, column_index=0, text="项目"),
                TableCell(row_index=0, column_index=1, text="期末余额"),
                TableCell(row_index=0, column_index=2, text="期末余额"),
                TableCell(row_index=1, column_index=0, text="应收账款"),
                TableCell(
                    row_index=1,
                    column_index=1,
                    text="100",
                    numeric_value="100",
                ),
                TableCell(
                    row_index=1,
                    column_index=2,
                    text="200",
                    numeric_value="200",
                ),
            ]
        }
    )
    document = ParsedDocument(
        document_id="document",
        document_version_id="version",
        document_type=DocumentType.FINANCIAL_REPORT,
        metadata={"issuer": "测试医药股份有限公司"},
        elements=[element],
    )
    mentions = EntityExtractAgent([PatternExtractor()]).extract(document)

    assert RelationExtractAgent().extract(document, mentions) == []


def test_non_metric_percentage_column_is_not_emitted_as_amount() -> None:
    element = make_element(
        document_version_id="version",
        element_type=ElementType.TABLE,
        reading_order=0,
        text="项目\t本报告期末\t本报告期末\n固定资产\t500\t10.5%",
        parser_name="mineru",
        parser_version="3.4.0",
    ).model_copy(
        update={
            "table_cells": [
                TableCell(row_index=0, column_index=0, text="项目"),
                TableCell(row_index=0, column_index=1, text="本报告期末"),
                TableCell(row_index=0, column_index=2, text="本报告期末"),
                TableCell(row_index=1, column_index=0, text="固定资产"),
                TableCell(row_index=1, column_index=1, text="500", numeric_value="500"),
                TableCell(
                    row_index=1,
                    column_index=2,
                    text="10.5%",
                    numeric_value="10.5",
                    unit="%",
                ),
            ]
        }
    )
    document = ParsedDocument(
        document_id="document",
        document_version_id="version",
        document_type=DocumentType.FINANCIAL_REPORT,
        metadata={"issuer": "测试医药股份有限公司"},
        elements=[element],
    )
    mentions = EntityExtractAgent([PatternExtractor()]).extract(document)
    assertions = RelationExtractAgent().extract(document, mentions)

    assert [item.object_value for item in assertions] == ["500"]
